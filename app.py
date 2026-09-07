import os
import re
import json
import asyncio
import uuid
import urllib.parse
from pathlib import Path
from typing import Optional, List, Dict, Any
from datetime import datetime

from fastapi import FastAPI, HTTPException, BackgroundTasks, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse, HTMLResponse
from pydantic import BaseModel

from database import DatabaseManager
from models import RawLead, ScoredLead, DigitalMaturityScore, WebsiteAuditResult, TechStackDetail, MultichannelSequence
from crm import CRMStage, CRMStageManager
from queue_manager import QueueManager
from auditor import WebsiteAuditor
from personalizer import OutreachPersonalizer
from proposals import ProposalGenerator
from report_generator import OpportunityReportGenerator
from lead_finder import GoogleMapsLeadFinder
from explainer import ScoreExplainer
from battlecard import SalesBattlecardGenerator
from scheduler import AutonomousScheduler
from territory_manager import TerritoryManager
from ai_qualifier import AIQualifier
from outreach_generator import OutreachGenerator
from timeline import OpportunityTimelineManager
from simulator import AICallSimulator

# Setup directories
BASE_DIR = Path(__file__).resolve().parent
DASHBOARD_DIR = BASE_DIR / "dashboard"
DASHBOARD_DIR.mkdir(exist_ok=True, parents=True)
if os.getenv("VERCEL"):
    OUTPUT_DIR = Path("/tmp/output")
else:
    OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True, parents=True)
(OUTPUT_DIR / "reports").mkdir(exist_ok=True, parents=True)
(OUTPUT_DIR / "proposals").mkdir(exist_ok=True, parents=True)
(OUTPUT_DIR / "screenshots").mkdir(exist_ok=True, parents=True)

def capture_screenshot_sync(url: str, save_path: Path) -> bool:
    target_url = url if (url.startswith("http://") or url.startswith("https://")) else f"https://{url}"
    try:
        from playwright.async_api import async_playwright
        async def _capture():
            async with async_playwright() as p:
                browser = await p.chromium.launch(headless=True)
                context = await browser.new_context(
                    viewport={"width": 1280, "height": 800},
                    user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
                )
                page = await context.new_page()
                try:
                    await page.goto(target_url, timeout=12000, wait_until="domcontentloaded")
                    await asyncio.sleep(1.0)
                    await page.screenshot(path=str(save_path), full_page=False, quality=80, type="jpeg")
                    return True
                except Exception:
                    return False
                finally:
                    await browser.close()
        return asyncio.run(_capture())
    except Exception:
        return False

app = FastAPI(title="Dental WhatsApp Growth & Lead Intelligence API", version="2.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Mount static directories
app.mount("/output", StaticFiles(directory=str(OUTPUT_DIR)), name="output")
app.mount("/static", StaticFiles(directory=str(DASHBOARD_DIR)), name="static")

# Shared Database Instance
db = DatabaseManager()

# Background job tracking
active_jobs: Dict[str, Dict[str, Any]] = {}

# --- Pydantic Request Models ---

class SearchRequest(BaseModel):
    query: str = "Dentists in Austin, TX"
    limit: int = 10
    generate_reports: bool = True
    auto_queue: bool = True

class StageUpdateRequest(BaseModel):
    stage: str
    notes: Optional[str] = "Moved via Visual Dashboard"

class NoteRequest(BaseModel):
    note: str
    author: Optional[str] = "Sales Rep"

class QueueActionRequest(BaseModel):
    notes: Optional[str] = None

class ProposalRequest(BaseModel):
    lead_id: Optional[str] = None
    url: Optional[str] = None
    name: Optional[str] = None

class CallOutcomeRequest(BaseModel):
    outcome: str  # INTERESTED, NOT_INTERESTED, ALREADY_HAS_AI, GATEKEEPER_BLOCKED, WON, LOST, CALLBACK_REQUESTED
    notes: Optional[str] = None
    duration_sec: int = 0

class NextTerritoryRequest(BaseModel):
    territory_id: Optional[str] = None
    limit: int = 8

class CallSimulationStartRequest(BaseModel):
    persona: Optional[str] = "GATEKEEPER_RECEPTIONIST"

class CallSimulationTurnRequest(BaseModel):
    persona: str = "GATEKEEPER_RECEPTIONIST"
    user_pitch: str
    history: List[Dict[str, str]] = []

class TimelineEventCreateRequest(BaseModel):
    event_type: str
    title: str
    description: Optional[str] = None
    actor: str = "SALES_REP"
    metadata: Optional[Dict[str, Any]] = None

class DaemonToggleRequest(BaseModel):
    enable: bool = True
    interval_hours: float = 6.0
    limit_per_cycle: int = 8

# --- Helper Functions ---

def format_clean_phone(raw_phone: Optional[str]) -> str:
    if not raw_phone:
        return ""
    digits = re.sub(r"\D", "", raw_phone)
    if len(digits) == 10:
        return f"1{digits}"
    elif len(digits) == 11 and digits.startswith("1"):
        return digits
    return digits

def generate_wa_pitch_and_link(lead_dict: Dict[str, Any]) -> Dict[str, str]:
    name = lead_dict.get("name", "Doctor")
    phone = lead_dict.get("phone", "")
    address = lead_dict.get("address", "Texas")
    rating = lead_dict.get("rating", 4.9)
    review_count = lead_dict.get("review_count", 50)
    score = lead_dict.get("opportunity_score", 75)
    annual_gap = lead_dict.get("annual_gap", 65000) or 65000
    monthly_rev = f"${int(annual_gap / 12):,}"

    if "dr." in name.lower() or "dr " in name.lower():
        salutation = f"Dr. {name.split(',')[0].replace('Dr.', '').replace('Dr', '').strip()}"
    else:
        salutation = f"{name} Team"

    clean_phone = format_clean_phone(phone)

    pitch = (
        f"Hi {salutation}! 👋 I was reviewing premier dental practices in {address} "
        f"and noticed {name}'s {rating}★ reputation across {review_count} patient reviews.\n\n"
        f"Quick question: Did you know ~42% of patient emergency searches & inquiries happen after 5 PM when the front desk is closed? "
        f"We estimate this is leaking roughly {monthly_rev}/month in uncaptured chair production.\n\n"
        f"We install a 24/7 WhatsApp Patient Automation system that:\n"
        f"• Instantly texts back missed callers within 5 seconds so they don't call another dentist\n"
        f"• Lets patients self-schedule exams & cleanings directly inside WhatsApp 24/7\n"
        f"• Fills last-minute cancellations with 98% open-rate broadcast recall\n\n"
        f"Would you be open to a 3-minute video showing how it links to your existing practice calendar?"
    )

    wa_link = f"https://wa.me/{clean_phone}?text={urllib.parse.quote_plus(pitch)}" if clean_phone else ""

    return {
        "salutation": salutation,
        "clean_phone": clean_phone,
        "pitch": pitch,
        "wa_link": wa_link
    }

# --- Routes ---

@app.get("/")
async def serve_dashboard():
    index_path = DASHBOARD_DIR / "index.html"
    if index_path.exists():
        return FileResponse(str(index_path))
    return HTMLResponse("<h1>Dashboard loading...</h1>")

@app.get("/api/stats")
async def get_stats():
    summary = db.get_crm_pipeline_summary()
    queue_pending = db.list_outreach_queue(status="PENDING_REVIEW")
    queue_approved = db.list_outreach_queue(status="APPROVED")
    queue_sent = db.list_outreach_queue(status="SENT")

    total_leads = sum(v["count"] for v in summary.values())
    total_leakage = sum(v["total_leakage"] for v in summary.values())

    return {
        "total_leads": total_leads,
        "total_leakage": total_leakage,
        "pipeline": summary,
        "queue": {
            "pending": len(queue_pending),
            "approved": len(queue_approved),
            "sent": len(queue_sent)
        }
    }

@app.get("/api/leads")
async def get_leads(
    stage: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = 100
):
    if stage and stage != "ALL":
        leads = db.get_leads_by_stage(stage)
    else:
        leads = db.list_all_leads()

    if search:
        search_lower = search.lower()
        leads = [l for l in leads if search_lower in (l.get("name") or "").lower() or search_lower in (l.get("phone") or "").lower() or search_lower in (l.get("website") or "").lower()]

    # Enhance leads with clean phone, revenue metrics, screenshots, and quick WhatsApp action
    results = []
    for l in leads[:limit]:
        wa_data = generate_wa_pitch_and_link(l)
        l_copy = dict(l)
        l_copy["clean_phone"] = wa_data["clean_phone"]
        l_copy["wa_link"] = wa_data["wa_link"]

        # Ensure numeric coordinates
        l_copy["latitude"] = float(l["latitude"]) if l.get("latitude") is not None else None
        l_copy["longitude"] = float(l["longitude"]) if l.get("longitude") is not None else None

        # Format missed revenue range
        rev_min = l.get("missed_rev_min") or l.get("audit_missed_rev_min") or 0
        rev_max = l.get("missed_rev_max") or l.get("audit_missed_rev_max") or 0
        if not rev_min and not rev_max:
            annual = l.get("annual_gap") or 0
            if annual > 0:
                rev_max = int(annual / 12)
                rev_min = int(rev_max * 0.6)
            else:
                rev_min = 3450
                rev_max = 6900
        l_copy["missed_rev_min"] = rev_min
        l_copy["missed_rev_max"] = rev_max
        l_copy["missed_rev_range"] = f"${rev_min:,.0f} - ${rev_max:,.0f}"

        # Screenshot URL
        lead_id = l.get("id")
        shot_path = l.get("screenshot_path")
        if shot_path and Path(shot_path).exists():
            l_copy["screenshot_url"] = f"/output/screenshots/{Path(shot_path).name}"
        elif lead_id and (OUTPUT_DIR / "screenshots" / f"{lead_id}.jpg").exists():
            l_copy["screenshot_url"] = f"/output/screenshots/{lead_id}.jpg"
        else:
            l_copy["screenshot_url"] = None

        results.append(l_copy)

    return results

@app.get("/api/leads/export")
async def export_leads(
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    stage: Optional[str] = None
):
    import io
    import csv
    from fastapi.responses import Response

    if stage and stage != "ALL":
        leads = db.get_leads_by_stage(stage)
    else:
        leads = db.list_all_leads()

    headers = [
        "Practice Name",
        "Doctor / Owner",
        "Decision Maker Role",
        "Category",
        "Rating",
        "Review Count",
        "Phone Number",
        "WhatsApp Direct Link",
        "Website",
        "Address",
        "Pipeline Stage",
        "Opportunity Score",
        "Est. Missed Monthly Revenue",
        "Est. Annual Revenue Leakage",
        "AI Strategic Why-Score Rationale",
        "Last Audit Date"
    ]

    rows = []
    for l in leads:
        wa_data = generate_wa_pitch_and_link(l)
        rev_min = l.get("missed_rev_min") or l.get("audit_missed_rev_min") or 0
        rev_max = l.get("missed_rev_max") or l.get("audit_missed_rev_max") or 0
        if not rev_min and not rev_max:
            annual = l.get("annual_gap") or 65000
            rev_max = int(annual / 12)
            rev_min = int(rev_max * 0.6)
        rev_str = f"${rev_min:,.0f} - ${rev_max:,.0f}/mo"
        annual_str = f"${(l.get('annual_gap') or 65000):,.0f}/yr"

        why_bullets = l.get("why_score_reasons") or []
        if isinstance(why_bullets, list):
            why_text = " • ".join(r.lstrip("•-* ").strip() for r in why_bullets if r)
        else:
            why_text = str(why_bullets)

        rows.append([
            l.get("name", ""),
            l.get("doctor_name") or "Not Identified",
            l.get("decision_maker_role") or "",
            l.get("category", "Dentist"),
            float(l.get("rating", 4.9)) if l.get("rating") else 4.9,
            int(l.get("review_count", 0)) if l.get("review_count") else 0,
            l.get("phone", ""),
            wa_data.get("wa_link", ""),
            l.get("website", ""),
            l.get("address", ""),
            l.get("stage", "FOUND"),
            int(l.get("opportunity_score", 0)),
            rev_str,
            annual_str,
            why_text,
            l.get("last_audit_date") or l.get("last_updated") or ""
        ])

    timestamp = datetime.now().strftime("%Y%m%d_%H%M")

    if format == "xlsx":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dental Prospects"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
        border_thin = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, row_data in enumerate(rows, start=2):
            ws.append(row_data)
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = border_thin
                cell.alignment = Alignment(vertical="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        ws.auto_filter.ref = ws.dimensions

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"dental_leads_export_{timestamp}.xlsx"

        return Response(
            content=buffer.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename=\"{filename}\"",
                "Access-Control-Expose-Headers": "Content-Disposition"
            }
        )
    else:
        output = io.StringIO()
        output.write('\ufeff')
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row)
        filename = f"dental_leads_export_{timestamp}.csv"
        return Response(
            content=output.getvalue().encode("utf-8-sig"),
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f"attachment; filename=\"{filename}\"",
                "Access-Control-Expose-Headers": "Content-Disposition"
            }
        )

@app.get("/api/leads/{lead_id}")
async def get_lead_details(lead_id: str):
    lead = db.get_lead(lead_id)
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")

    findings = db.get_lead_findings(lead_id)
    notes = db.get_lead_notes(lead_id)
    history = db.get_stage_history(lead_id)
    wa_data = generate_wa_pitch_and_link(lead)

    # Fetch technologies
    with db._get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM technologies WHERE lead_id = ?", (lead_id,))
        techs = [dict(r) for r in cursor.fetchall()]

        # Latest audit details
        cursor.execute("SELECT * FROM audits WHERE lead_id = ? ORDER BY timestamp DESC LIMIT 1", (lead_id,))
        audit = cursor.fetchone()
        audit_dict = dict(audit) if audit else {}

    # Check for generated proposal
    clean_name = "".join(c for c in lead.get("name", "") if c.isalnum() or c == " ").strip().replace(" ", "_")
    proposal_file = OUTPUT_DIR / "proposals" / f"{clean_name}_growth_proposal.html"
    proposal_url = f"/output/proposals/{proposal_file.name}" if proposal_file.exists() else None

    # Check for generated PDF report
    pdf_file = OUTPUT_DIR / "reports" / f"{clean_name}_audit.pdf"
    pdf_url = f"/output/reports/{pdf_file.name}" if pdf_file.exists() else None

    return {
        "lead": lead,
        "audit": audit_dict,
        "findings": findings,
        "technologies": techs,
        "notes": notes,
        "history": history,
        "whatsapp": wa_data,
        "proposal_url": proposal_url,
        "pdf_url": pdf_url
    }

@app.get("/api/leads/{lead_id}/battlecard")
async def get_lead_battlecard(lead_id: str):
    lead = db.get_lead(lead_id)
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    with db._get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM audits WHERE lead_id = ? ORDER BY timestamp DESC LIMIT 1", (lead_id,))
        audit = cursor.fetchone()
        if audit:
            audit_dict = dict(audit)
            lead["detected_chatbot_name"] = audit_dict.get("detected_chatbot_name")
            lead["detected_booking_tool"] = audit_dict.get("detected_booking_tool")

    battlecard = SalesBattlecardGenerator.generate_battlecard(lead)
    return battlecard

@app.get("/api/briefing/today")
async def get_today_briefing():
    all_leads = db.list_all_leads()
    sorted_leads = sorted(all_leads, key=lambda l: (l.get("opportunity_score") or 0), reverse=True)

    top_targets = []
    total_monthly_leakage = 0
    doctors_found = 0

    for l in sorted_leads[:10]:
        rev_min = l.get("missed_rev_min") or l.get("audit_missed_rev_min") or 3500
        rev_max = l.get("missed_rev_max") or l.get("audit_missed_rev_max") or 6800
        total_monthly_leakage += rev_max
        doc = l.get("doctor_name")
        if doc and doc != "Not Identified":
            doctors_found += 1

        wa_data = generate_wa_pitch_and_link(l)
        top_targets.append({
            "id": l.get("id"),
            "name": l.get("name"),
            "doctor_name": doc or "Primary Dentist",
            "decision_maker_role": l.get("decision_maker_role") or "Practice Owner",
            "rating": float(l.get("rating", 4.9)) if l.get("rating") else 4.9,
            "reviews": int(l.get("review_count", 0)) if l.get("review_count") else 0,
            "phone": l.get("phone", ""),
            "website": l.get("website", ""),
            "address": l.get("address", ""),
            "opportunity_score": int(l.get("opportunity_score", 75)),
            "stage": l.get("stage", "FOUND"),
            "missed_rev_min": rev_min,
            "missed_rev_max": rev_max,
            "annual_gap": (rev_min + rev_max) // 2 * 12,
            "wa_link": wa_data.get("wa_link", ""),
            "wa_pitch": wa_data.get("wa_pitch", ""),
            "why_score_reasons": l.get("why_score_reasons", [])
        })

    total_annual_arr = total_monthly_leakage * 12

    return {
        "briefing_date": datetime.now().strftime("%A, %B %d, %Y"),
        "total_targets_ready": len(top_targets),
        "doctors_identified_count": doctors_found,
        "combined_monthly_missed_revenue": total_monthly_leakage,
        "total_annual_arr_opportunity": total_annual_arr,
        "targets": top_targets
    }

@app.post("/api/briefing/run-morning")
async def run_morning_harvest(city: Optional[str] = "Austin, TX", limit: int = 5):
    result = await AutonomousScheduler.run_morning_cycle(
        cities=[city],
        limit_per_city=limit,
        headless=True,
        db=db
    )
    return {"status": "completed", "summary": result}

# --- Autonomous Territory & Sales Employee Endpoints ---

@app.get("/api/territories")
async def get_territories():
    TerritoryManager.initialize_territories(db)
    territories = db.get_all_territories()
    pending = sum(1 for t in territories if t.get("status") == "PENDING")
    in_progress = sum(1 for t in territories if t.get("status") == "IN_PROGRESS")
    completed = sum(1 for t in territories if t.get("status") == "COMPLETED")
    next_target = TerritoryManager.get_next_target_territory(db)
    return {
        "total_territories": len(territories),
        "pending": pending,
        "in_progress": in_progress,
        "completed": completed,
        "next_target": next_target,
        "territories": territories
    }

@app.post("/api/territories/next-cycle")
async def trigger_next_territory_cycle(req: NextTerritoryRequest, background_tasks: BackgroundTasks):
    summary = await AutonomousScheduler.run_autonomous_cycle(
        territory_id=req.territory_id,
        limit=req.limit,
        headless=True,
        db=db
    )
    return {"status": "completed", "summary": summary}

@app.get("/api/queue/todays-calls")
async def get_todays_calls(min_probability: int = 70):
    leads = db.get_todays_calls(min_probability=min_probability)
    results = []
    for l in leads:
        wa_data = generate_wa_pitch_and_link(l)
        l_copy = dict(l)
        l_copy["clean_phone"] = wa_data["clean_phone"]
        l_copy["wa_link"] = wa_data["wa_link"]

        rev_min = l.get("missed_rev_min") or l.get("audit_missed_rev_min") or 3450
        rev_max = l.get("missed_rev_max") or l.get("audit_missed_rev_max") or 6900
        l_copy["missed_rev_range"] = f"${rev_min:,.0f} - ${rev_max:,.0f}"

        shot_path = l.get("screenshot_path")
        lead_id = l.get("id")
        if shot_path and Path(shot_path).exists():
            l_copy["screenshot_url"] = f"/output/screenshots/{Path(shot_path).name}"
        elif lead_id and (OUTPUT_DIR / "screenshots" / f"{lead_id}.jpg").exists():
            l_copy["screenshot_url"] = f"/output/screenshots/{lead_id}.jpg"
        else:
            l_copy["screenshot_url"] = None

        results.append(l_copy)
    return results

@app.post("/api/leads/{lead_id}/call-outcome")
async def log_call_outcome(lead_id: str, req: CallOutcomeRequest):
    log_id = db.log_call_outcome(
        lead_id=lead_id,
        outcome=req.outcome,
        rep_notes=req.notes,
        duration_sec=req.duration_sec
    )
    lead = db.get_lead(lead_id)
    return {
        "status": "success",
        "log_id": log_id,
        "lead_id": lead_id,
        "recorded_outcome": req.outcome,
        "new_crm_stage": lead.get("stage") if lead else "UNKNOWN"
    }

@app.get("/api/leads/{lead_id}/outreach")
async def get_lead_outreach(lead_id: str):
    lead = db.get_lead(lead_id)
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    
    rev_min = lead.get("missed_rev_min") or 3450
    rev_max = lead.get("missed_rev_max") or 6900
    monthly_leakage = (rev_min + rev_max) // 2

    outreach = OutreachGenerator.generate(
        lead_dict=dict(lead),
        doctor_name=lead.get("doctor_name"),
        monthly_leakage=monthly_leakage
    )
    return outreach.model_dump()

@app.get("/api/learning/insights")
async def get_learning_insights():
    return db.get_conversion_intelligence()

# --- Opportunity Timeline Endpoints ---

@app.get("/api/leads/{lead_id}/timeline")
async def get_lead_timeline(lead_id: str):
    events = db.get_lead_timeline(lead_id)
    return {"lead_id": lead_id, "events": events, "count": len(events)}

@app.post("/api/leads/{lead_id}/timeline")
async def add_lead_timeline_event(lead_id: str, req: TimelineEventCreateRequest):
    event_id = db.log_timeline_event(
        lead_id=lead_id,
        event_type=req.event_type,
        title=req.title,
        description=req.description,
        actor=req.actor,
        metadata=req.metadata
    )
    return {"status": "success", "event_id": event_id, "lead_id": lead_id}

# --- AI Call Roleplay Simulator Endpoints ---

@app.post("/api/leads/{lead_id}/simulate-call/start")
async def start_call_simulation(lead_id: str, req: Optional[CallSimulationStartRequest] = None):
    lead = db.get_lead(lead_id)
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    persona = req.persona if req else "GATEKEEPER_RECEPTIONIST"
    session_data = AICallSimulator.start_session(dict(lead), persona=persona)
    return session_data

@app.post("/api/leads/{lead_id}/simulate-call/turn")
async def simulate_call_turn(lead_id: str, req: CallSimulationTurnRequest):
    lead = db.get_lead(lead_id)
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    turn_res = AICallSimulator.simulate_turn(
        lead_dict=dict(lead),
        persona=req.persona,
        conversation_history=req.history,
        user_pitch=req.user_pitch
    )
    # Log practiced simulation event to lead timeline
    try:
        db.log_timeline_event(
            lead_id=lead_id,
            event_type="SIMULATION_PRACTICED",
            title=f"Call Sparring: {req.persona} (Score: {turn_res.evaluation.overall_score}/10)",
            description=f"Turn evaluated. Status: {turn_res.status}. Feedback: {turn_res.evaluation.tactical_feedback}",
            actor="SALES_REP",
            metadata={
                "persona": req.persona,
                "overall_score": turn_res.evaluation.overall_score,
                "hook_score": turn_res.evaluation.hook_score,
                "value_score": turn_res.evaluation.value_score,
                "control_score": turn_res.evaluation.control_score,
                "status": turn_res.status
            }
        )
    except Exception:
        pass
    return turn_res.model_dump()

# --- Expected Value ($EV) Prioritized Queue ---

@app.get("/api/queue/top-ev")
async def get_top_ev_queue(min_probability: int = 50, limit: int = 50):
    leads = db.get_top_ev_queue(min_probability=min_probability, limit=limit)
    results = []
    for l in leads:
        wa_data = generate_wa_pitch_and_link(l)
        l_copy = dict(l)
        l_copy["clean_phone"] = wa_data["clean_phone"]
        l_copy["wa_link"] = wa_data["wa_link"]

        rev_min = l.get("missed_rev_min") or l.get("audit_missed_rev_min") or 3450
        rev_max = l.get("missed_rev_max") or l.get("audit_missed_rev_max") or 6900
        l_copy["missed_rev_range"] = f"${rev_min:,.0f} - ${rev_max:,.0f}"

        shot_path = l.get("screenshot_path")
        lead_id = l.get("id")
        if shot_path and Path(shot_path).exists():
            l_copy["screenshot_url"] = f"/output/screenshots/{Path(shot_path).name}"
        elif lead_id and (OUTPUT_DIR / "screenshots" / f"{lead_id}.jpg").exists():
            l_copy["screenshot_url"] = f"/output/screenshots/{lead_id}.jpg"
        else:
            l_copy["screenshot_url"] = None

        results.append(l_copy)
    return results

# --- 24/7 Autonomous Daemon Endpoints ---

@app.get("/api/scheduler/daemon")
async def get_daemon_status():
    return AutonomousScheduler.get_daemon_status()

@app.post("/api/scheduler/daemon/toggle")
async def toggle_daemon(req: DaemonToggleRequest):
    if req.enable:
        status = await AutonomousScheduler.start_daemon(
            interval_hours=req.interval_hours,
            limit_per_cycle=req.limit_per_cycle,
            headless=True,
            db=db
        )
    else:
        status = AutonomousScheduler.stop_daemon()
    return status

@app.post("/api/leads/{lead_id}/stage")
async def update_lead_stage(lead_id: str, req: StageUpdateRequest):
    success = db.update_lead_stage(lead_id, new_stage=req.stage, notes=req.notes)
    if not success:
        raise HTTPException(status_code=400, detail="Failed to update stage")
    return {"status": "success", "lead_id": lead_id, "new_stage": req.stage}

@app.post("/api/leads/{lead_id}/note")
async def add_lead_note(lead_id: str, req: NoteRequest):
    note_id = db.add_lead_note(lead_id, note_text=req.note, author=req.author or "Sales Rep")
    return {"status": "success", "note_id": note_id}

# --- Queue Management Endpoints ---

@app.get("/api/queue")
async def list_queue(status: Optional[str] = "PENDING_REVIEW"):
    filter_status = None if status == "ALL" else status
    items = db.list_outreach_queue(status=filter_status)
    results = []
    for item in items:
        item_copy = dict(item)
        clean_phone = format_clean_phone(item.get("phone"))
        item_copy["clean_phone"] = clean_phone
        item_copy["wa_link"] = f"https://wa.me/{clean_phone}?text={urllib.parse.quote_plus(item.get('body', ''))}" if clean_phone else ""
        results.append(item_copy)
    return results

@app.post("/api/queue/{queue_id}/approve")
async def approve_queue_item(queue_id: str, req: QueueActionRequest):
    success = QueueManager.approve_item(queue_id, review_notes=req.notes or "Approved via Visual Dashboard", db=db)
    if not success:
        raise HTTPException(status_code=400, detail="Failed to approve item")
    return {"status": "success", "queue_id": queue_id}

@app.post("/api/queue/{queue_id}/reject")
async def reject_queue_item(queue_id: str, req: QueueActionRequest):
    success = QueueManager.reject_item(queue_id, reason=req.notes or "Rejected via Visual Dashboard", db=db)
    if not success:
        raise HTTPException(status_code=400, detail="Failed to reject item")
    return {"status": "success", "queue_id": queue_id}

@app.post("/api/queue/dispatch")
async def dispatch_approved_queue():
    dispatched = QueueManager.dispatch_approved(db=db)
    return {"status": "success", "count": len(dispatched), "items": dispatched}

# --- Proposal Generator Endpoint ---

@app.post("/api/proposal/generate")
async def generate_proposal(req: ProposalRequest):
    name = req.name or "Premier Dental Clinic"
    url = req.url or "https://example.com"

    if req.lead_id:
        lead = db.get_lead(req.lead_id)
        if lead:
            name = lead.get("name", name)
            url = lead.get("website", url)

    clean_name = "".join(c for c in name if c.isalnum() or c == " ").strip().replace(" ", "_")
    filepath = OUTPUT_DIR / "proposals" / f"{clean_name}_growth_proposal.html"

    # Run proposal engine
    raw = RawLead(name=name, website=url, phone="555-0100")
    scored = ScoredLead(
        raw_lead=raw,
        audit=WebsiteAuditResult(),
        opportunity_score=85,
        estimated_missed_revenue_annual=71550,
        estimated_missed_revenue_monthly_max=5962
    )
    proposal_path = ProposalGenerator.generate_proposal(
        scored_lead=scored,
        agency_name="WhatsApp Growth Partners for Dentists"
    )

    return {
        "status": "success",
        "proposal_path": str(proposal_path),
        "proposal_url": f"/output/proposals/{proposal_path.name}",
        "tier1_setup": 1500,
        "tier1_monthly": 350,
        "tier2_setup": 2500,
        "tier2_monthly": 650
    }

# --- Live Search Background Task & Progress ---

async def run_search_background_task(job_id: str, query: str, limit: int, generate_reports: bool, auto_queue: bool):
    job = active_jobs[job_id]
    job["status"] = "running"
    job["message"] = f"Connecting to Google Maps for: {query}..."
    job["progress"] = 5
    job["clinics_found"] = 0
    job["websites_analyzed"] = 0
    job["high_opportunity_count"] = 0
    job["logs"].append(f"[{datetime.now().strftime('%H:%M:%S')}] Started Google Maps scraper for '{query}'...")

    try:
        finder = GoogleMapsLeadFinder(headless=True)

        async def on_lead_found(raw_lead, count, total_target):
            job["clinics_found"] = count
            job["message"] = f"Found clinic ({count}/{limit}): {raw_lead.name}..."
            job["logs"].append(f"[{datetime.now().strftime('%H:%M:%S')}] Found practice: {raw_lead.name} ({raw_lead.phone or 'checking website'})...")
            pct = int(5 + (count / max(limit, 1)) * 30)
        def is_existing_fn(name, website, phone):
            return db.is_lead_existing(name=name, website=website, phone=phone)

        async def on_lead_skipped(name, reason):
            job["logs"].append(f"[{datetime.now().strftime('%H:%M:%S')}] ⏩ Skipping '{name}' ({reason})")

        raw_leads = await finder.search(
            query=query,
            limit=limit,
            on_lead_found=on_lead_found,
            is_existing_fn=is_existing_fn,
            on_lead_skipped=on_lead_skipped
        )
        job["total_found"] = len(raw_leads)
        job["message"] = f"Found {len(raw_leads)} practices. Inspecting websites & generating intelligence..."
        job["progress"] = 35
        job["logs"].append(f"[{datetime.now().strftime('%H:%M:%S')}] Discovery complete. Auditing digital maturity for {len(raw_leads)} practices...")

        auditor = WebsiteAuditor()
        personalizer = OutreachPersonalizer()
        explainer = ScoreExplainer()

        completed_leads = []
        step_increment = 60 / max(len(raw_leads), 1)

        for i, raw in enumerate(raw_leads):
            job["message"] = f"Auditing ({i+1}/{len(raw_leads)}): {raw.name}..."
            job["logs"].append(f"[{datetime.now().strftime('%H:%M:%S')}] Auditing {raw.name} (Tech stack & missed revenue)...")

            scored_lead = await auditor.audit_lead(raw)
            scored_lead = await personalizer.personalize_lead(scored_lead)
            lead_id = db._make_lead_id(raw.name, raw.website)

            # Capture screenshot if website exists
            screenshot_path = None
            if raw.website:
                shot_file = OUTPUT_DIR / "screenshots" / f"{lead_id}.jpg"
                if not shot_file.exists():
                    job["logs"].append(f"[{datetime.now().strftime('%H:%M:%S')}] 📸 Capturing homepage screenshot for {raw.name}...")
                    try:
                        success = await capture_website_screenshot(raw.website, shot_file)
                        if success:
                            screenshot_path = str(shot_file)
                    except Exception:
                        pass
                else:
                    screenshot_path = str(shot_file)

            # Generate Gemini explanation for score
            job["logs"].append(f"[{datetime.now().strftime('%H:%M:%S')}] 🤖 Generating AI Opportunity Insights for {raw.name} (Score: {scored_lead.opportunity_score})...")
            try:
                why_reasons = await asyncio.to_thread(explainer.explain_score, scored_lead)
            except Exception:
                why_reasons = explainer.fallback_reasons(scored_lead)
            scored_lead.why_score_reasons = why_reasons
            if screenshot_path:
                scored_lead.screenshot_path = screenshot_path

            # Update DB with screenshot and AI explanation
            db.update_lead_visuals(lead_id, screenshot_path=screenshot_path, why_reasons=why_reasons)

            # Auto-generate PDF report if requested
            if generate_reports and scored_lead.opportunity_score >= 50:
                try:
                    await OpportunityReportGenerator.generate_pdf(scored_lead)
                except Exception:
                    pass

            # Auto-stage in review queue if high value
            if auto_queue and (scored_lead.opportunity_score >= 60 or (scored_lead.deal_priority and "TIER 1" in scored_lead.deal_priority.tier.value)):
                try:
                    QueueManager.stage_for_review(scored_lead, db=db)
                except Exception:
                    pass

            # Update progress counters
            job["websites_analyzed"] = i + 1
            if scored_lead.opportunity_score >= 60:
                job["high_opportunity_count"] += 1

            wa_data = generate_wa_pitch_and_link({
                "name": raw.name,
                "phone": raw.phone,
                "address": raw.address,
                "rating": raw.rating,
                "review_count": raw.review_count,
                "opportunity_score": scored_lead.opportunity_score,
                "annual_gap": scored_lead.estimated_missed_revenue_annual
            })

            rev_min = scored_lead.estimated_missed_revenue_monthly_min or int((scored_lead.estimated_missed_revenue_annual or 60000) / 12 * 0.6)
            rev_max = scored_lead.estimated_missed_revenue_monthly_max or int((scored_lead.estimated_missed_revenue_annual or 60000) / 12)
            screenshot_url = f"/output/screenshots/{Path(screenshot_path).name}" if (screenshot_path and Path(screenshot_path).exists()) else None

            completed_leads.append({
                "id": lead_id,
                "name": raw.name,
                "phone": raw.phone,
                "clean_phone": wa_data["clean_phone"],
                "wa_link": wa_data["wa_link"],
                "website": raw.website,
                "rating": raw.rating,
                "review_count": raw.review_count,
                "address": raw.address,
                "latitude": raw.latitude,
                "longitude": raw.longitude,
                "opportunity_score": scored_lead.opportunity_score,
                "score": scored_lead.opportunity_score,
                "missed_rev_min": rev_min,
                "missed_rev_max": rev_max,
                "missed_rev_range": f"${rev_min:,.0f} - ${rev_max:,.0f}",
                "annual_gap": scored_lead.estimated_missed_revenue_annual,
                "annual_loss": scored_lead.estimated_missed_revenue_annual,
                "screenshot_url": screenshot_url,
                "why_score_reasons": why_reasons,
                "stage": "AUDITED"
            })

            job["progress"] = int(35 + (i + 1) * step_increment)
            job["discovered_leads"] = completed_leads

        job["status"] = "completed"
        job["message"] = f"Search complete! Discovered and analyzed {len(completed_leads)} dental practices."
        job["progress"] = 100
        job["logs"].append(f"[{datetime.now().strftime('%H:%M:%S')}] Discovery & intelligence run complete. {len(completed_leads)} practices ready.")

    except Exception as e:
        job["status"] = "error"
        job["message"] = f"Search encountered an issue: {str(e)}"
        job["logs"].append(f"[{datetime.now().strftime('%H:%M:%S')}] ERROR: {str(e)}")

@app.post("/api/search/start")
async def start_search(req: SearchRequest, background_tasks: BackgroundTasks):
    job_id = f"job_{datetime.now().strftime('%Y%m%d%H%M%S')}_{str(uuid.uuid4())[:6]}"
    active_jobs[job_id] = {
        "job_id": job_id,
        "query": req.query,
        "status": "pending",
        "progress": 0,
        "message": "Initializing discovery...",
        "total_found": 0,
        "clinics_found": 0,
        "websites_analyzed": 0,
        "high_opportunity_count": 0,
        "logs": [f"[{datetime.now().strftime('%H:%M:%S')}] Initializing search for '{req.query}'..."],
        "discovered_leads": []
    }

    background_tasks.add_task(
        run_search_background_task,
        job_id=job_id,
        query=req.query,
        limit=req.limit,
        generate_reports=req.generate_reports,
        auto_queue=req.auto_queue
    )

    return {"status": "started", "job_id": job_id}

@app.get("/api/search/status/{job_id}")
async def get_search_status(job_id: str):
    if job_id not in active_jobs:
        raise HTTPException(status_code=404, detail="Job ID not found")
    return active_jobs[job_id]

if __name__ == "__main__":
    import webbrowser
    import uvicorn

    print("=" * 60)
    print("  Enterprise Dental WhatsApp Intelligence Platform")
    print("  Local Server starting on: http://127.0.0.1:8000")
    print("=" * 60)

    # Open browser automatically after 1.5 seconds
    def open_browser():
        webbrowser.open("http://127.0.0.1:8000")

    import threading
    threading.Timer(1.5, open_browser).start()

    uvicorn.run(app, host="127.0.0.1", port=8000, log_level="info")
